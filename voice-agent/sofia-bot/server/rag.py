#
# HQ Lead Engine — RAG (knowledge base) for voice assistants.
#
# index(assistant_dir): read docs in <dir>/knowledge/, extract text, chunk,
#   embed (OpenAI), and save vectors to <dir>/knowledge_index.json.
# search(assistant_dir, query, k): embed the query and return the top-k most
#   relevant chunks (cosine similarity). Called by the engine's
#   query_knowledge_base tool at call time.
#
import json
import os
import re

import numpy as np

EMBED_MODEL = "text-embedding-3-small"
INDEX_FILE = "knowledge_index.json"
KB_SUBDIR = "knowledge"


def _client():
    from openai import OpenAI

    return OpenAI(api_key=os.getenv("OPENAI_API_KEY"))


def extract_text(path: str) -> str:
    ext = path.lower().rsplit(".", 1)[-1] if "." in path else ""
    try:
        if ext in ("txt", "md", "csv"):
            with open(path, encoding="utf-8", errors="ignore") as f:
                return f.read()
        if ext == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            return "\n".join((pg.extract_text() or "") for pg in reader.pages)
        if ext in ("docx",):
            import docx

            d = docx.Document(path)
            return "\n".join(p.text for p in d.paragraphs)
        if ext in ("xlsx", "xlsm"):
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        out.append(" | ".join(cells))
            return "\n".join(out)
    except Exception as e:  # noqa: BLE001
        print(f"[rag] extract failed for {path}: {e}")
    return ""


def chunk_text(text: str, size: int = 900, overlap: int = 150) -> list[str]:
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return []
    chunks, i = [], 0
    while i < len(text):
        chunks.append(text[i : i + size])
        i += size - overlap
    return [c for c in chunks if c.strip()]


def embed_texts(texts: list[str]) -> list[list[float]]:
    if not texts:
        return []
    resp = _client().embeddings.create(model=EMBED_MODEL, input=texts)
    return [d.embedding for d in resp.data]


def index_path(assistant_dir: str) -> str:
    return os.path.join(assistant_dir, INDEX_FILE)


def index(assistant_dir: str) -> dict:
    """(Re)build the vector index from all docs in <assistant_dir>/knowledge/."""
    kb = os.path.join(assistant_dir, KB_SUBDIR)
    records: list[dict] = []
    files = 0
    if os.path.isdir(kb):
        for name in sorted(os.listdir(kb)):
            p = os.path.join(kb, name)
            if not os.path.isfile(p):
                continue
            text = extract_text(p)
            if not text.strip():
                continue
            files += 1
            for c in chunk_text(text):
                records.append({"text": c, "source": name})

    # embed in batches of 100
    for start in range(0, len(records), 100):
        batch = [r["text"] for r in records[start : start + 100]]
        for rec, emb in zip(records[start : start + 100], embed_texts(batch)):
            rec["embedding"] = emb

    with open(index_path(assistant_dir), "w", encoding="utf-8") as f:
        json.dump({"chunks": records}, f)
    return {"files": files, "chunks": len(records)}


def has_index(assistant_dir: str) -> bool:
    return os.path.isfile(index_path(assistant_dir))


def index_stats(assistant_dir: str) -> dict:
    if not has_index(assistant_dir):
        return {"indexed": False, "chunks": 0}
    try:
        with open(index_path(assistant_dir), encoding="utf-8") as f:
            data = json.load(f)
        return {"indexed": True, "chunks": len(data.get("chunks", []))}
    except Exception:
        return {"indexed": False, "chunks": 0}


def search(assistant_dir: str, query: str, k: int = 4) -> list[str]:
    """Return the top-k most relevant chunk texts for the query."""
    if not has_index(assistant_dir) or not query.strip():
        return []
    with open(index_path(assistant_dir), encoding="utf-8") as f:
        chunks = json.load(f).get("chunks", [])
    chunks = [c for c in chunks if c.get("embedding")]
    if not chunks:
        return []
    q = np.array(embed_texts([query])[0], dtype=np.float32)
    M = np.array([c["embedding"] for c in chunks], dtype=np.float32)
    sims = (M @ q) / (np.linalg.norm(M, axis=1) * (np.linalg.norm(q) + 1e-9) + 1e-9)
    top = np.argsort(-sims)[:k]
    return [chunks[i]["text"] for i in top]
